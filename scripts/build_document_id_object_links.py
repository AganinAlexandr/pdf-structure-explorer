#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import unicodedata
import zlib
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


DEFAULT_WORKBOOK = Path(r"E:\output\pdf-structure-explorer\Elements_PDF.xlsx")
DEFAULT_EXPORTS_ROOT = Path(r"E:\output\pdf-structure-explorer\exports")
DEFAULT_OBJECT_MAP = Path(r"E:\commons\pdf-structure-explorer\reference\simplified_objects_section_map.xlsx")
DEFAULT_OUTPUT_STEM = Path(r"E:\output\pdf-structure-explorer\Elements_PDF_document_links")

TRUTHY_VALUES = {"true", "1", "истина", "да", "yes"}

LINK_COLUMNS = [
    "document_id",
    "file_name",
    "file_crc32",
    "file_size_bytes",
    "export_document_id",
    "object_code",
    "object_name",
    "document_role",
    "section_type",
    "section_filter",
    "subsection_type",
    "subsection_filter",
    "source_catalog_group",
    "partial_bundle",
    "match_found",
    "match_count",
    "match_mode",
]

PIVOT_OUTPUT_COLUMNS = [
    "document_id",
    "frames",
    "images",
    "lines",
    "other_vector",
    "tables",
    "text",
    "total",
    "file_name",
    "file_crc32",
    "file_size_bytes",
    "object_code",
    "object_name",
    "document_role",
    "section_type",
    "section_filter",
    "partial_bundle",
    "match_found",
    "match_count",
    "match_mode",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build document_id -> object/section links for Elements_PDF workbook.")
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK))
    parser.add_argument("--exports-root", default=str(DEFAULT_EXPORTS_ROOT))
    parser.add_argument("--object-map", default=str(DEFAULT_OBJECT_MAP))
    parser.add_argument("--output-stem", default=str(DEFAULT_OUTPUT_STEM))
    return parser.parse_args()


def normalize_text(value: str) -> str:
    return unicodedata.normalize("NFC", str(value or "")).strip()


def normalize_file_name(value: str) -> str:
    return normalize_text(value).casefold()


def normalize_crc32(value: str) -> str:
    return normalize_text(value).upper()


def normalize_size_bytes(value: str) -> str:
    text = normalize_text(value)
    if not text:
        return ""
    try:
        return str(int(float(text)))
    except ValueError:
        return text


def normalize_stored_path(value: str) -> Path | None:
    text = normalize_text(value)
    if not text:
        return None
    return Path(text.replace("/", "\\"))


def is_truthy(value: str) -> bool:
    return normalize_text(value).casefold() in TRUTHY_VALUES


def compute_crc32(path: Path) -> str:
    crc = 0
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            crc = zlib.crc32(chunk, crc)
    return f"{crc & 0xFFFFFFFF:08X}"


def load_xlsx_rows(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    header = [str(v) if v is not None else "" for v in values[0]]
    rows: list[dict[str, str]] = []
    for row_values in values[1:]:
        row = {}
        for idx, column in enumerate(header):
            value = row_values[idx] if idx < len(row_values) else ""
            row[column] = "" if value is None else str(value)
        rows.append(row)
    return rows, header


def load_object_map(path: Path) -> tuple[
    dict[tuple[str, str], list[dict[str, str]]],
    dict[str, list[dict[str, str]]],
]:
    rows, _ = load_xlsx_rows(path)
    grouped_by_name_crc: dict[tuple[str, str], list[dict[str, str]]] = {}
    grouped_by_crc: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        key = (
            normalize_file_name(row.get("file_name", "")),
            normalize_crc32(row.get("file_crc32", "")),
        )
        if not key[1]:
            continue
        if key[0]:
            grouped_by_name_crc.setdefault(key, []).append(row)
        grouped_by_crc.setdefault(key[1], []).append(row)
    return grouped_by_name_crc, grouped_by_crc


def load_export_documents(exports_root: Path) -> dict[str, dict[str, str]]:
    documents = {}
    for folder in exports_root.glob("doc_*"):
        doc_csv = folder / "documents.csv"
        if not doc_csv.exists():
            continue
        with doc_csv.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            try:
                row = next(reader)
            except StopIteration:
                continue
        row["partial_bundle"] = "true" if row.get("parse_status", "") not in ("", "parsed") else "false"
        if not normalize_crc32(row.get("file_crc32", "")):
            stored_path = normalize_stored_path(row.get("file_path", ""))
            if stored_path and stored_path.exists():
                row["file_crc32"] = compute_crc32(stored_path)
        documents[row["document_id"]] = row
    return documents


def extract_workbook_document_ids(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    pivot = workbook.worksheets[0]

    nonempty_rows = []
    for index, row in enumerate(pivot.iter_rows(values_only=True), start=1):
        values = [value for value in row if value not in (None, "")]
        if values:
            nonempty_rows.append((index, values))

    pivot_rows: list[dict[str, str]] = []
    workbook_ids = []
    for _, values in nonempty_rows:
        document_id = normalize_text(values[0])
        if not document_id.startswith("doc_"):
            continue
        workbook_ids.append(document_id)
        pivot_rows.append({
            "document_id": document_id,
            "frames": str(values[1]) if len(values) > 1 else "",
            "images": str(values[2]) if len(values) > 2 else "",
            "lines": str(values[3]) if len(values) > 3 else "",
            "other_vector": str(values[4]) if len(values) > 4 else "",
            "tables": str(values[5]) if len(values) > 5 else "",
            "text": str(values[6]) if len(values) > 6 else "",
            "total": str(values[7]) if len(values) > 7 else "",
        })

    for sheet in workbook.worksheets[1:]:
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            continue
        header_list = [str(value) if value is not None else "" for value in header]
        if "document_id" not in header_list:
            continue
        doc_idx = header_list.index("document_id")
        for row in rows:
            value = row[doc_idx] if doc_idx < len(row) else None
            if value:
                workbook_ids.append(str(value))

    unique_ids = sorted(set(workbook_ids))
    return unique_ids, pivot_rows


def choose_preferred_match(matches: list[dict[str, str]]) -> dict[str, str]:
    return next((row for row in matches if is_truthy(row.get("is_preferred_source", ""))), matches[0])


def choose_object_match(
    doc_row: dict[str, str],
    object_map_by_name_crc: dict[tuple[str, str], list[dict[str, str]]],
    object_map_by_crc: dict[str, list[dict[str, str]]],
):
    key = (
        normalize_file_name(doc_row.get("file_name", "")),
        normalize_crc32(doc_row.get("file_crc32", "")),
    )
    matches = object_map_by_name_crc.get(key, [])
    if matches:
        return choose_preferred_match(matches), len(matches), "file_name+file_crc32"

    crc = key[1]
    if not crc:
        return None, 0, ""

    crc_matches = object_map_by_crc.get(crc, [])
    if not crc_matches:
        return None, 0, ""

    size = normalize_size_bytes(doc_row.get("file_size_bytes", ""))
    if size:
        size_matches = [
            row for row in crc_matches
            if normalize_size_bytes(row.get("file_size_bytes", "")) == size
        ]
        if size_matches:
            return choose_preferred_match(size_matches), len(size_matches), "crc32_only"

    return choose_preferred_match(crc_matches), len(crc_matches), "crc32_only"


def build_links(document_ids: list[str], export_docs: dict[str, dict[str, str]], object_map_by_name_crc, object_map_by_crc):
    stats = Counter()
    rows = []
    for document_id in document_ids:
        export_row = export_docs.get(document_id)
        if not export_row:
            stats["missing_export_doc"] += 1
            rows.append({
                "document_id": document_id,
                "file_name": "",
                "file_crc32": "",
                "file_size_bytes": "",
                "export_document_id": "",
                "object_code": "",
                "object_name": "",
                "document_role": "",
                "section_type": "",
                "section_filter": "",
                "subsection_type": "",
                "subsection_filter": "",
                "source_catalog_group": "",
                "partial_bundle": "false",
                "match_found": "false",
                "match_count": "0",
                "match_mode": "missing_export_doc",
            })
            continue

        linked, match_count, match_mode = choose_object_match(export_row, object_map_by_name_crc, object_map_by_crc)
        if linked:
            stats["matched"] += 1
            stats[f"matched_{match_mode}"] += 1
            rows.append({
                "document_id": document_id,
                "file_name": export_row.get("file_name", ""),
                "file_crc32": export_row.get("file_crc32", ""),
                "file_size_bytes": export_row.get("file_size_bytes", ""),
                "export_document_id": export_row.get("document_id", ""),
                "object_code": linked.get("object_code", ""),
                "object_name": linked.get("object_name", ""),
                "document_role": linked.get("document_role", ""),
                "section_type": linked.get("section_type", ""),
                "section_filter": linked.get("section_filter", ""),
                "subsection_type": linked.get("subsection_type", ""),
                "subsection_filter": linked.get("subsection_filter", ""),
                "source_catalog_group": linked.get("source_catalog_group", ""),
                "partial_bundle": export_row.get("partial_bundle", "false"),
                "match_found": "true",
                "match_count": str(match_count),
                "match_mode": match_mode,
            })
        else:
            stats["unmatched_object_map"] += 1
            rows.append({
                "document_id": document_id,
                "file_name": export_row.get("file_name", ""),
                "file_crc32": export_row.get("file_crc32", ""),
                "file_size_bytes": export_row.get("file_size_bytes", ""),
                "export_document_id": export_row.get("document_id", ""),
                "object_code": "",
                "object_name": "",
                "document_role": "",
                "section_type": "",
                "section_filter": "",
                "subsection_type": "",
                "subsection_filter": "",
                "source_catalog_group": "",
                "partial_bundle": export_row.get("partial_bundle", "false"),
                "match_found": "false",
                "match_count": "0",
                "match_mode": "no_object_map_match",
            })
    return rows, stats


def build_link_index(link_rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    return {row["document_id"]: row for row in link_rows}


def build_pivot_linked_rows(pivot_rows: list[dict[str, str]], link_index: dict[str, dict[str, str]]):
    rows = []
    for row in pivot_rows:
        linked = link_index.get(row["document_id"], {})
        rows.append({
            "document_id": row["document_id"],
            "frames": row["frames"],
            "images": row["images"],
            "lines": row["lines"],
            "other_vector": row["other_vector"],
            "tables": row["tables"],
            "text": row["text"],
            "total": row["total"],
            "file_name": linked.get("file_name", ""),
            "file_crc32": linked.get("file_crc32", ""),
            "file_size_bytes": linked.get("file_size_bytes", ""),
            "object_code": linked.get("object_code", ""),
            "object_name": linked.get("object_name", ""),
            "document_role": linked.get("document_role", ""),
            "section_type": linked.get("section_type", ""),
            "section_filter": linked.get("section_filter", ""),
            "partial_bundle": linked.get("partial_bundle", "false"),
            "match_found": linked.get("match_found", "false"),
            "match_count": linked.get("match_count", "0"),
            "match_mode": linked.get("match_mode", ""),
        })
    return rows


def write_csv(path: Path, rows: list[dict[str, str]], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, link_rows: list[dict[str, str]], pivot_rows: list[dict[str, str]]) -> None:
    workbook = Workbook()
    sheet_links = workbook.active
    sheet_links.title = "document_links"
    sheet_links.append(LINK_COLUMNS)
    for row in link_rows:
        sheet_links.append([row.get(column, "") for column in LINK_COLUMNS])

    sheet_pivot = workbook.create_sheet("pivot_linked")
    sheet_pivot.append(PIVOT_OUTPUT_COLUMNS)
    for row in pivot_rows:
        sheet_pivot.append([row.get(column, "") for column in PIVOT_OUTPUT_COLUMNS])

    for sheet in (sheet_links, sheet_pivot):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    workbook.save(path)


def main() -> int:
    args = parse_args()
    workbook_path = Path(args.workbook)
    exports_root = Path(args.exports_root)
    object_map_path = Path(args.object_map)
    output_stem = Path(args.output_stem)

    document_ids, pivot_rows = extract_workbook_document_ids(workbook_path)
    export_docs = load_export_documents(exports_root)
    object_map_by_name_crc, object_map_by_crc = load_object_map(object_map_path)
    link_rows, stats = build_links(document_ids, export_docs, object_map_by_name_crc, object_map_by_crc)
    pivot_linked_rows = build_pivot_linked_rows(pivot_rows, build_link_index(link_rows))

    csv_path = output_stem.with_suffix(".csv")
    xlsx_path = output_stem.with_suffix(".xlsx")
    pivot_csv_path = output_stem.with_name(f"{output_stem.name}_pivot.csv")

    write_csv(csv_path, link_rows, LINK_COLUMNS)
    write_csv(pivot_csv_path, pivot_linked_rows, PIVOT_OUTPUT_COLUMNS)
    write_xlsx(xlsx_path, link_rows, pivot_linked_rows)

    print(f"Workbook document_ids: {len(document_ids)}")
    print(f"Matched to export docs: {len(document_ids) - stats['missing_export_doc']}")
    print(f"Matched to object map: {stats['matched']}")
    print(f"Matched by file_name+file_crc32: {stats['matched_file_name+file_crc32']}")
    print(f"Matched by crc32 only (renamed copies): {stats['matched_crc32_only']}")
    print(f"Missing export docs: {stats['missing_export_doc']}")
    print(f"No object-map match: {stats['unmatched_object_map']}")
    print(f"Links CSV: {csv_path}")
    print(f"Pivot CSV: {pivot_csv_path}")
    print(f"Output XLSX: {xlsx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
