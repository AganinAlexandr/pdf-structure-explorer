#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


DEFAULT_MAP = Path(r"E:\commons\pdf-structure-explorer\reference\simplified_objects_section_map.xlsx")
MAP_FIELDS = [
    "object_code",
    "object_name",
    "document_role",
    "section_type",
    "section_filter",
    "subsection_type",
    "subsection_filter",
    "source_catalog_group",
]
OUTPUT_FIELDS = [
    "linked_object_code",
    "linked_object_name",
    "linked_document_role",
    "linked_section_type",
    "linked_section_filter",
    "linked_subsection_type",
    "linked_subsection_filter",
    "linked_source_catalog_group",
    "linked_match_mode",
    "linked_match_count",
    "linked_match_found",
]
PATH_CANDIDATES = ["full_path", "file_path", "path"]
FILE_NAME_CANDIDATES = ["file_name", "filename", "file"]
CRC_CANDIDATES = ["file_crc32", "crc32"]
SIZE_CANDIDATES = ["file_size_bytes", "size_bytes", "file_size"]
TRUTHY_VALUES = {"true", "1", "истина", "да", "yes"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Merge a PDF results table with object/section metadata.",
    )
    parser.add_argument("--input", required=True, help="Input results table (.csv or .xlsx)")
    parser.add_argument(
        "--map",
        default=str(DEFAULT_MAP),
        help=f"Object/section map file (.csv or .xlsx). Default: {DEFAULT_MAP}",
    )
    parser.add_argument(
        "--output-stem",
        help="Optional output stem without extension. By default uses <input>_linked next to the input file.",
    )
    return parser.parse_args()


def normalize_text(value: str) -> str:
    return unicodedata.normalize("NFC", str(value or "")).strip()


def normalize_path(value: str) -> str:
    text = normalize_text(value).replace("/", "\\")
    while "\\\\" in text:
        text = text.replace("\\\\", "\\")
    return text.casefold()


def normalize_crc32(value: str) -> str:
    return normalize_text(value).upper()


def normalize_size_bytes(value: str) -> str:
    text = normalize_text(value)
    if not text:
        return ""
    try:
        return str(int(float(text)))
    except ValueError:
        return text


def normalize_file_name(value: str) -> str:
    return normalize_text(value).casefold()


def is_truthy(value: str) -> bool:
    return normalize_text(value).casefold() in TRUTHY_VALUES


def find_column(columns: Iterable[str], candidates: list[str]) -> str | None:
    normalized = {column.casefold(): column for column in columns}
    for candidate in candidates:
        if candidate.casefold() in normalized:
            return normalized[candidate.casefold()]
    return None


def load_csv(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
        return rows, list(reader.fieldnames or [])


def load_xlsx(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return [], []
    header = [str(value) if value is not None else "" for value in values[0]]
    rows: list[dict[str, str]] = []
    for row_values in values[1:]:
        row = {}
        for index, column in enumerate(header):
            value = row_values[index] if index < len(row_values) else ""
            row[column] = "" if value is None else str(value)
        rows.append(row)
    return rows, header


def load_table(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return load_csv(path)
    if suffix == ".xlsx":
        return load_xlsx(path)
    raise ValueError(f"Unsupported table format: {path.suffix}")


def choose_preferred_match(matches: list[dict[str, str]]) -> dict[str, str]:
    return next((item for item in matches if is_truthy(item.get("is_preferred_source", ""))), matches[0])


def choose_match(row: dict[str, str], map_by_path, map_by_name_crc, map_by_crc, path_col, file_name_col, crc_col, size_col):
    if path_col:
        key = normalize_path(row.get(path_col, ""))
        if key and key in map_by_path:
            matches = map_by_path[key]
            return matches[0], "full_path", len(matches)
    if file_name_col and crc_col:
        key = (
            normalize_file_name(row.get(file_name_col, "")),
            normalize_crc32(row.get(crc_col, "")),
        )
        if key[0] and key[1] and key in map_by_name_crc:
            matches = map_by_name_crc[key]
            return choose_preferred_match(matches), "file_name+file_crc32", len(matches)
    if crc_col:
        crc = normalize_crc32(row.get(crc_col, ""))
        if crc and crc in map_by_crc:
            matches = map_by_crc[crc]
            size = normalize_size_bytes(row.get(size_col, "")) if size_col else ""
            if size:
                size_matches = [
                    item for item in matches
                    if normalize_size_bytes(item.get("file_size_bytes", "")) == size
                ]
                if size_matches:
                    return choose_preferred_match(size_matches), "crc32_only", len(size_matches)
            return choose_preferred_match(matches), "crc32_only", len(matches)
    return None, "", 0


def build_indices(rows: list[dict[str, str]], columns: list[str]):
    path_col = find_column(columns, PATH_CANDIDATES)
    file_name_col = find_column(columns, FILE_NAME_CANDIDATES)
    crc_col = find_column(columns, CRC_CANDIDATES)
    size_col = find_column(columns, SIZE_CANDIDATES)

    by_path = defaultdict(list)
    by_name_crc = defaultdict(list)
    by_crc = defaultdict(list)

    for row in rows:
        if path_col:
            key = normalize_path(row.get(path_col, ""))
            if key:
                by_path[key].append(row)
        if file_name_col and crc_col:
            key = (
                normalize_file_name(row.get(file_name_col, "")),
                normalize_crc32(row.get(crc_col, "")),
            )
            if key[0] and key[1]:
                by_name_crc[key].append(row)
            if key[1]:
                by_crc[key[1]].append(row)
    return by_path, by_name_crc, by_crc, path_col, file_name_col, crc_col, size_col


def merge_rows(input_rows, input_columns, map_rows, map_columns):
    map_by_path, map_by_name_crc, map_by_crc, _map_path_col, _map_file_name_col, _map_crc_col, _map_size_col = build_indices(map_rows, map_columns)
    input_path_col = find_column(input_columns, PATH_CANDIDATES)
    input_file_name_col = find_column(input_columns, FILE_NAME_CANDIDATES)
    input_crc_col = find_column(input_columns, CRC_CANDIDATES)
    input_size_col = find_column(input_columns, SIZE_CANDIDATES)

    stats = Counter()
    merged = []
    for row in input_rows:
        linked, match_mode, match_count = choose_match(
            row,
            map_by_path,
            map_by_name_crc,
            map_by_crc,
            input_path_col,
            input_file_name_col,
            input_crc_col,
            input_size_col,
        )
        out = dict(row)
        if linked:
            for field in MAP_FIELDS:
                out[f"linked_{field}"] = linked.get(field, "")
            out["linked_match_mode"] = match_mode
            out["linked_match_count"] = str(match_count)
            out["linked_match_found"] = "true"
            stats[f"matched_{match_mode}"] += 1
        else:
            for field in MAP_FIELDS:
                out[f"linked_{field}"] = ""
            out["linked_match_mode"] = ""
            out["linked_match_count"] = "0"
            out["linked_match_found"] = "false"
            stats["unmatched"] += 1
        merged.append(out)

    output_columns = list(input_columns)
    for field in OUTPUT_FIELDS:
        if field not in output_columns:
            output_columns.append(field)
    return merged, output_columns, stats


def write_csv(path: Path, rows: list[dict[str, str]], columns: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, rows: list[dict[str, str]], columns: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "linked_results"
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])

    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for index, column in enumerate(columns, start=1):
        width = 18
        if "path" in column:
            width = 72
        elif "file_name" in column:
            width = 42
        elif "object_name" in column:
            width = 32
        elif "section" in column:
            width = 18
        elif "notes" in column:
            width = 28
        sheet.column_dimensions[_column_letter(index)].width = width

    workbook.save(path)


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def main() -> int:
    args = parse_args()
    input_path = Path(args.input)
    map_path = Path(args.map)
    output_stem = Path(args.output_stem) if args.output_stem else input_path.with_name(f"{input_path.stem}_linked")

    input_rows, input_columns = load_table(input_path)
    map_rows, map_columns = load_table(map_path)
    merged_rows, output_columns, stats = merge_rows(input_rows, input_columns, map_rows, map_columns)

    output_csv = output_stem.with_suffix(".csv")
    output_xlsx = output_stem.with_suffix(".xlsx")
    write_csv(output_csv, merged_rows, output_columns)
    write_xlsx(output_xlsx, merged_rows, output_columns)

    print(f"Input rows: {len(input_rows)}")
    print(f"Map rows: {len(map_rows)}")
    print(f"Matched by full_path: {stats['matched_full_path']}")
    print(f"Matched by file_name+file_crc32: {stats['matched_file_name+file_crc32']}")
    print(f"Matched by crc32 only (renamed copies): {stats['matched_crc32_only']}")
    print(f"Unmatched: {stats['unmatched']}")
    print(f"Output CSV: {output_csv}")
    print(f"Output XLSX: {output_xlsx}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
