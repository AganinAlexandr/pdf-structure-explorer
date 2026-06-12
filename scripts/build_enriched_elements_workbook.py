#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


DEFAULT_WORKBOOK = Path(r"E:\output\pdf-structure-explorer\Elements_PDF.xlsx")
DEFAULT_LINKS = Path(r"E:\output\pdf-structure-explorer\Elements_PDF_document_links.csv")
DEFAULT_PIVOT_LINKS = Path(r"E:\output\pdf-structure-explorer\Elements_PDF_document_links_pivot.csv")
DEFAULT_OUTPUT = Path(r"E:\output\pdf-structure-explorer\Elements_PDF_enriched.xlsx")

LINK_FIELDS = [
    "object_code",
    "object_name",
    "document_role",
    "section_type",
    "section_filter",
    "subsection_type",
    "subsection_filter",
    "source_catalog_group",
    "partial_bundle",
    "match_found",
    "match_count",
    "match_mode",
]

SECTION_DOCUMENT_COLUMNS = [
    "document_id",
    "file_name",
    "file_crc32",
    "object_code",
    "object_name",
    "section_type",
    "section_filter",
    "source_catalog_group",
    "partial_bundle",
    "match_mode",
]

OBJECT_SECTION_SUMMARY_COLUMNS = [
    "object_code",
    "object_name",
    "section_type",
    "section_filter",
    "document_count",
    "frames_total",
    "images_total",
    "lines_total",
    "other_vector_total",
    "tables_total",
    "text_total",
    "elements_total",
]

OBJECT_SUMMARY_COLUMNS = [
    "object_code",
    "object_name",
    "document_count",
    "section_document_count",
    "non_section_document_count",
    "distinct_section_count",
    "frames_total",
    "images_total",
    "lines_total",
    "other_vector_total",
    "tables_total",
    "text_total",
    "elements_total",
]

SECTION_SUMMARY_COLUMNS = [
    "section_type",
    "section_filter",
    "document_count",
    "object_count",
    "frames_total",
    "images_total",
    "lines_total",
    "other_vector_total",
    "tables_total",
    "text_total",
    "elements_total",
]

UNMATCHED_COLUMNS = [
    "document_id",
    "file_name",
    "file_crc32",
    "match_mode",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create an enriched workbook with object and section metadata.")
    parser.add_argument("--workbook", default=str(DEFAULT_WORKBOOK))
    parser.add_argument("--links", default=str(DEFAULT_LINKS))
    parser.add_argument("--pivot-links", default=str(DEFAULT_PIVOT_LINKS))
    parser.add_argument("--output", default=str(DEFAULT_OUTPUT))
    return parser.parse_args()


def normalize_text(value) -> str:
    return "" if value is None else str(value)


def as_int(value: str) -> int:
    text = normalize_text(value)
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def load_csv_rows(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        rows = list(reader)
        return rows, list(reader.fieldnames or [])


def build_link_index(path: Path) -> tuple[dict[str, dict[str, str]], list[dict[str, str]], list[str]]:
    rows, columns = load_csv_rows(path)
    index = {row["document_id"]: row for row in rows if row.get("document_id")}
    return index, rows, columns


def build_section_documents(link_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    out = []
    for row in link_rows:
        if row.get("match_found") != "true" or row.get("document_role") != "section":
            continue
        out.append({column: row.get(column, "") for column in SECTION_DOCUMENT_COLUMNS})
    out.sort(key=lambda row: (row["object_code"], row["section_type"], row["file_name"].casefold()))
    return out


def build_object_section_summary(pivot_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    grouped = {}
    for row in pivot_rows:
        if row.get("match_found") != "true" or row.get("document_role") != "section":
            continue
        key = (row.get("object_code", ""), row.get("object_name", ""), row.get("section_type", ""), row.get("section_filter", ""))
        agg = grouped.setdefault(key, {
            "object_code": key[0],
            "object_name": key[1],
            "section_type": key[2],
            "section_filter": key[3],
            "document_count": 0,
            "frames_total": 0,
            "images_total": 0,
            "lines_total": 0,
            "other_vector_total": 0,
            "tables_total": 0,
            "text_total": 0,
            "elements_total": 0,
        })
        agg["document_count"] += 1
        agg["frames_total"] += as_int(row.get("frames", ""))
        agg["images_total"] += as_int(row.get("images", ""))
        agg["lines_total"] += as_int(row.get("lines", ""))
        agg["other_vector_total"] += as_int(row.get("other_vector", ""))
        agg["tables_total"] += as_int(row.get("tables", ""))
        agg["text_total"] += as_int(row.get("text", ""))
        agg["elements_total"] += as_int(row.get("total", ""))
    rows = list(grouped.values())
    rows.sort(key=lambda row: (row["object_code"], row["section_type"]))
    return [{column: str(row[column]) if isinstance(row[column], int) else row[column] for column in OBJECT_SECTION_SUMMARY_COLUMNS} for row in rows]


def build_object_summary(pivot_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    grouped = {}
    for row in pivot_rows:
        if row.get("match_found") != "true":
            continue
        key = (row.get("object_code", ""), row.get("object_name", ""))
        agg = grouped.setdefault(key, {
            "object_code": key[0],
            "object_name": key[1],
            "document_count": 0,
            "section_document_count": 0,
            "non_section_document_count": 0,
            "sections": set(),
            "frames_total": 0,
            "images_total": 0,
            "lines_total": 0,
            "other_vector_total": 0,
            "tables_total": 0,
            "text_total": 0,
            "elements_total": 0,
        })
        agg["document_count"] += 1
        if row.get("document_role") == "section":
            agg["section_document_count"] += 1
            if row.get("section_type"):
                agg["sections"].add(row["section_type"])
        else:
            agg["non_section_document_count"] += 1
        agg["frames_total"] += as_int(row.get("frames", ""))
        agg["images_total"] += as_int(row.get("images", ""))
        agg["lines_total"] += as_int(row.get("lines", ""))
        agg["other_vector_total"] += as_int(row.get("other_vector", ""))
        agg["tables_total"] += as_int(row.get("tables", ""))
        agg["text_total"] += as_int(row.get("text", ""))
        agg["elements_total"] += as_int(row.get("total", ""))
    rows = []
    for row in grouped.values():
        rows.append({
            "object_code": row["object_code"],
            "object_name": row["object_name"],
            "document_count": str(row["document_count"]),
            "section_document_count": str(row["section_document_count"]),
            "non_section_document_count": str(row["non_section_document_count"]),
            "distinct_section_count": str(len(row["sections"])),
            "frames_total": str(row["frames_total"]),
            "images_total": str(row["images_total"]),
            "lines_total": str(row["lines_total"]),
            "other_vector_total": str(row["other_vector_total"]),
            "tables_total": str(row["tables_total"]),
            "text_total": str(row["text_total"]),
            "elements_total": str(row["elements_total"]),
        })
    rows.sort(key=lambda row: row["object_code"])
    return rows


def build_section_summary(pivot_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    grouped = {}
    for row in pivot_rows:
        if row.get("match_found") != "true" or row.get("document_role") != "section" or not row.get("section_type"):
            continue
        key = (row.get("section_type", ""), row.get("section_filter", ""))
        agg = grouped.setdefault(key, {
            "section_type": key[0],
            "section_filter": key[1],
            "document_count": 0,
            "objects": set(),
            "frames_total": 0,
            "images_total": 0,
            "lines_total": 0,
            "other_vector_total": 0,
            "tables_total": 0,
            "text_total": 0,
            "elements_total": 0,
        })
        agg["document_count"] += 1
        if row.get("object_code"):
            agg["objects"].add(row["object_code"])
        agg["frames_total"] += as_int(row.get("frames", ""))
        agg["images_total"] += as_int(row.get("images", ""))
        agg["lines_total"] += as_int(row.get("lines", ""))
        agg["other_vector_total"] += as_int(row.get("other_vector", ""))
        agg["tables_total"] += as_int(row.get("tables", ""))
        agg["text_total"] += as_int(row.get("text", ""))
        agg["elements_total"] += as_int(row.get("total", ""))
    rows = []
    for row in grouped.values():
        rows.append({
            "section_type": row["section_type"],
            "section_filter": row["section_filter"],
            "document_count": str(row["document_count"]),
            "object_count": str(len(row["objects"])),
            "frames_total": str(row["frames_total"]),
            "images_total": str(row["images_total"]),
            "lines_total": str(row["lines_total"]),
            "other_vector_total": str(row["other_vector_total"]),
            "tables_total": str(row["tables_total"]),
            "text_total": str(row["text_total"]),
            "elements_total": str(row["elements_total"]),
        })
    rows.sort(key=lambda row: row["section_type"])
    return rows


def build_unmatched_rows(link_rows: list[dict[str, str]]) -> list[dict[str, str]]:
    rows = []
    for row in link_rows:
        if row.get("match_found") == "true":
            continue
        rows.append({column: row.get(column, "") for column in UNMATCHED_COLUMNS})
    rows.sort(key=lambda row: row["document_id"])
    return rows


def autofit(sheet, widths: dict[int, int]) -> None:
    for column_index, width in widths.items():
        sheet.column_dimensions[_column_letter(column_index)].width = width


def _column_letter(index: int) -> str:
    result = ""
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def append_table_sheet(workbook: Workbook, title: str, rows: list[dict[str, str]], columns: list[str]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {}
    for idx, column in enumerate(columns, start=1):
        width = 18
        if "path" in column:
            width = 72
        elif "file_name" in column:
            width = 42
        elif "object_name" in column:
            width = 32
        elif "section" in column:
            width = 18
        widths[idx] = width
    autofit(sheet, widths)


def enrich_sheet_rows(header: list[str], data_rows: list[list[str]], link_index: dict[str, dict[str, str]]) -> tuple[list[str], list[list[str]]]:
    if "document_id" not in header:
        return header, data_rows

    doc_idx = header.index("document_id")
    enriched_header = header + [f"linked_{field}" for field in LINK_FIELDS]
    enriched_rows = []
    for row in data_rows:
        doc_id = row[doc_idx] if doc_idx < len(row) else ""
        linked = link_index.get(doc_id, {})
        enriched_row = list(row)
        for field in LINK_FIELDS:
            enriched_row.append(linked.get(field, ""))
        enriched_rows.append(enriched_row)
    return enriched_header, enriched_rows


def main() -> int:
    args = parse_args()
    source_path = Path(args.workbook)
    links_path = Path(args.links)
    pivot_links_path = Path(args.pivot_links)
    output_path = Path(args.output)

    link_index, link_rows, link_columns = build_link_index(links_path)
    pivot_rows, pivot_columns = load_csv_rows(pivot_links_path)

    source_wb = load_workbook(source_path, read_only=True, data_only=True)
    out_wb = Workbook()
    out_wb.remove(out_wb.active)

    append_table_sheet(out_wb, "00_document_links", link_rows, link_columns)
    append_table_sheet(out_wb, "00_pivot_linked", pivot_rows, pivot_columns)
    append_table_sheet(out_wb, "01_section_docs", build_section_documents(link_rows), SECTION_DOCUMENT_COLUMNS)
    append_table_sheet(out_wb, "02_obj_section_sum", build_object_section_summary(pivot_rows), OBJECT_SECTION_SUMMARY_COLUMNS)
    append_table_sheet(out_wb, "03_object_sum", build_object_summary(pivot_rows), OBJECT_SUMMARY_COLUMNS)
    append_table_sheet(out_wb, "04_section_sum", build_section_summary(pivot_rows), SECTION_SUMMARY_COLUMNS)
    append_table_sheet(out_wb, "05_unmatched_docs", build_unmatched_rows(link_rows), UNMATCHED_COLUMNS)

    for sheet in source_wb.worksheets[1:]:
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            continue
        header = [normalize_text(value) for value in header_row]
        data_rows = [[normalize_text(value) for value in row] for row in rows_iter]
        enriched_header, enriched_rows = enrich_sheet_rows(header, data_rows, link_index)

        out_sheet = out_wb.create_sheet(f"{sheet.title}_linked"[:31])
        out_sheet.append(enriched_header)
        for row in enriched_rows:
            out_sheet.append(row)
        for cell in out_sheet[1]:
            cell.font = Font(bold=True)
        out_sheet.freeze_panes = "A2"
        out_sheet.auto_filter.ref = out_sheet.dimensions

        widths = {}
        for idx, column in enumerate(enriched_header, start=1):
            width = 16
            if "document_id" == column:
                width = 18
            elif "file_name" in column:
                width = 42
            elif "object_name" in column:
                width = 32
            elif "section" in column:
                width = 18
            elif "path" in column:
                width = 60
            widths[idx] = width
        autofit(out_sheet, widths)

    out_wb.save(output_path)
    print(f"Output workbook: {output_path}")
    print(f"Document links rows: {len(link_rows)}")
    print(f"Pivot linked rows: {len(pivot_rows)}")
    print(f"Linked sheets created: {len(source_wb.worksheets) - 1}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
